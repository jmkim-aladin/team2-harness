#!/usr/bin/env python3
"""YouTrack 루트 이슈 하위 트리 → WBS·간트 엑셀(.xlsx) + 선택적 HTML 간트 시제품.

루트 이슈 하나에서 Subtask 링크(parent for)를 재귀로 따라가 프로젝트 경계와 무관하게
하위 트리 전체를 모은다. 트리 깊이·유형(Epic/Feature/Task …)에 의존하지 않는다.
YouTrack은 GET만 호출한다(관찰 전용). 토큰 값은 헤더에만 쓰고 어디에도 출력하지 않는다.
엑셀과 HTML은 같은 계산 결과(진척률·rollup·점검)를 쓴다 — 두 산출물의 숫자가 같아야 비교가 된다.

시트: WBS / 간트(일 또는 주 단위 칸 색칠, 행 그룹) / 팀별 요약 / 점검 / 메타

일정 근사:
- 계획 = Sprints(월 단위 version, 예 2026.09) 최소월 1일 ~ 최대월 말일 (planned_dates 한 곳)
- 실제 시작 = State(현지화 이름 `상태` 포함)가 처음 In Progress가 된 시각
- 실제 종료 = 이슈 resolved 시각
- 하위 집계(rollup) = 상위 행의 자기 일정이 없거나 하위가 더 넓을 때 하위 범위를 빗금으로 표시

사용법:
    python3 tools/export_wbs.py DEV2-5237 --out ~/Downloads/wbs-DEV2-5237.xlsx
    python3 tools/export_wbs.py DEV2-5237 --out ~/Downloads/wbs.xlsx --html ~/Downloads/gantt.html
    python3 tools/export_wbs.py DEV2-5237 --out /tmp/wbs.xlsx --gantt-unit week --today 2026-10-07

토큰: env YOUTRACK_TOKEN 우선, 없으면 `python3 tools/cred.py get youtrack-token`.
base URL: env YOUTRACK_BASE_URL (기본 https://aladincommunication.youtrack.cloud).
"""
from __future__ import annotations

import argparse
import calendar
import datetime as dt
import html
import json
import os
import re
import subprocess
import sys
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional, Tuple

REPO = Path(__file__).resolve().parents[1]
HTML_TEMPLATE = Path(__file__).resolve().with_name("export_wbs_gantt.html")
KST = dt.timezone(dt.timedelta(hours=9))
DEFAULT_BASE_URL = "https://aladincommunication.youtrack.cloud"
TOKEN_CRED_NAME = "youtrack-token"

SUBTASK = "Subtask"
DEPEND = "Depend"
RELATES = "Relates"
STATE_FIELD_NAMES = {"State", "상태"}  # 활동 API는 필드 이름을 현지화해서 준다
IN_PROGRESS = "In Progress"
NOT_STARTED_STATES = {"ToBe", "Submitted", "Open", "Backlog"}
CANCEL_STATES = {"Won't fix", "Duplicate", "Obsolete", "Incomplete"}
TEAM_NAMES = {"DEV2": "개발2팀", "DEVSY": "기획팀", "DEVDG": "디자인팀"}
CHIP_KEYS = ("done_without_progress", "no_sprints", "prefix_variant", "overdue", "dependency_violation")

# 엑셀 WBS 열 ↔ HTML 인라인 데이터 키 — 두 산출물이 같은 열을 보여주는지 테스트가 이 표로 대조한다
WBS_LABELS = ["WBS", "레벨", "티켓", "제목", "유형", "프로젝트(팀)", "담당자", "상태", "구분", "Sprints",
              "예측", "소요 시간", "SP", "계획 시작", "계획 종료", "실제 시작", "실제 종료", "선행 티켓",
              "진척률(취소 제외)", "leaf 종료/대상"]
HTML_KEYS = {"WBS": "wbs", "레벨": "lv", "티켓": "id", "제목": "title", "유형": "type", "프로젝트(팀)": "team",
             "담당자": "asg", "상태": "state", "구분": "b", "Sprints": "spr", "예측": "est", "소요 시간": "spent",
             "SP": "sp", "계획 시작": "ps", "계획 종료": "pe", "실제 시작": "s0", "실제 종료": "s1",
             "선행 티켓": "deps", "진척률(취소 제외)": "prog", "leaf 종료/대상": "tot"}

# 유형 칸 색 — 엑셀 WBS '유형' 칸 채움과 HTML 표 '유형' 배지가 이 표 하나를 쓴다(HTML은 payload "type_colors").
# 간트 막대 색(계획/실제/지연/취소)·행 배경(레벨)은 다른 의미라 유형 칸에만 칠한다. 엑셀은 라이트 값을 쓴다.
TYPE_COLORS: Dict[str, Dict[str, str]] = {
    "epic": {"bg": "EBDDFB", "ink": "5B2A9E", "bg_dark": "4B3277", "ink_dark": "EEE2FF"},  # 보라
    "feature": {"bg": "C9EFEA", "ink": "0A5C55", "bg_dark": "164F4A", "ink_dark": "BDF2E9"},  # 청록
    "task": {"bg": "DCDFE4", "ink": "374151", "bg_dark": "434A55", "ink_dark": "E5E8EC"},  # 회색
    "bug": {"bg": "FBD8D8", "ink": "A1161F", "bg_dark": "6E2229", "ink_dark": "FFD9D9"},  # 빨강
    "other": {"bg": "F4F4F2", "ink": "6E6E6E", "bg_dark": "2A2F36", "ink_dark": "A4ACB6"},  # 그 밖·빈 값
}
TYPE_COLOR_FALLBACK = "other"

SPRINT_RE = re.compile(r"^\s*(\d{4})\.(\d{1,2})(?!\d)")
PREFIX_RE = re.compile(r"^\s*\[([^\]]+)\]")

ISSUE_FIELDS = (
    "idReadable,summary,resolved,project(shortName),"
    "customFields(name,value(name,login,fullName,presentation,minutes)),"
    "links(direction,linkType(name),issues(idReadable,resolved))"
)
ACTIVITY_FIELDS = "timestamp,field(name),added(name),removed(name)"
ISSUE_BATCH = 40
ACTIVITY_TOP = 2000
RETRIES = 2

BUCKET_DONE = "종료"
BUCKET_DOING = "진행"
BUCKET_TODO = "미착수"
BUCKET_CANCEL = "취소"

Range = Tuple[Optional[dt.date], Optional[dt.date]]


# ---------------------------------------------------------------------------
# 정규화


def id_key(issue_id: str) -> Tuple[str, int, str]:
    proj, _, num = issue_id.rpartition("-")
    return (proj, int(num) if num.isdigit() else 0, issue_id)


def ms_to_date(ms: Optional[int]) -> Optional[dt.date]:
    if not ms:
        return None
    return dt.datetime.fromtimestamp(ms / 1000, KST).date()


def _name(value: Any) -> Optional[str]:
    if isinstance(value, dict):
        return value.get("fullName") or value.get("name") or value.get("login") or value.get("presentation")
    if value is None:
        return None
    return str(value)


def _names(value: Any) -> List[str]:
    if value is None:
        return []
    items = value if isinstance(value, list) else [value]
    return [n for n in (_name(v) for v in items) if n]


def _presentation(value: Any) -> Optional[str]:
    if isinstance(value, dict):
        return value.get("presentation")
    return None


def normalize_issue(raw: Dict[str, Any]) -> Dict[str, Any]:
    """API 이슈 JSON → 내부 레코드. 링크는 Subtask/Depend 방향별로 분리한다."""
    iid = raw.get("idReadable") or ""
    fields = {cf.get("name"): cf.get("value") for cf in raw.get("customFields") or []}
    children: List[str] = []
    parents: List[str] = []
    depends: List[Tuple[str, bool]] = []
    required_for: List[Tuple[str, bool]] = []
    relates: List[Tuple[str, bool]] = []
    for link in raw.get("links") or []:
        link_type = (link.get("linkType") or {}).get("name")
        direction = link.get("direction")
        linked = [i for i in link.get("issues") or [] if i.get("idReadable")]
        if link_type == SUBTASK and direction == "OUTWARD":  # parent for
            children.extend(i["idReadable"] for i in linked)
        elif link_type == SUBTASK and direction == "INWARD":  # subtask of
            parents.extend(i["idReadable"] for i in linked)
        elif link_type == DEPEND and direction == "INWARD":  # depends on
            depends.extend((i["idReadable"], bool(i.get("resolved"))) for i in linked)
        elif link_type == DEPEND and direction == "OUTWARD":  # is required for
            required_for.extend((i["idReadable"], bool(i.get("resolved"))) for i in linked)
        elif link_type == RELATES:
            relates.extend((i["idReadable"], bool(i.get("resolved"))) for i in linked)
    sp = fields.get("Story points")
    return {
        "id": iid,
        "summary": raw.get("summary") or "",
        "project": (raw.get("project") or {}).get("shortName") or iid.rpartition("-")[0],
        "resolved": ms_to_date(raw.get("resolved")),
        "type": _name(fields.get("Type")),
        "state": _name(fields.get("State")),
        "assignee": _name(fields.get("Assignee")),
        "team": _name(fields.get("담당팀")),
        "sprints": _names(fields.get("Sprints")),
        "estimate": _presentation(fields.get("예측")),
        "spent": _presentation(fields.get("소요 시간")),
        "sp": sp if isinstance(sp, (int, float)) else None,
        "children": children,
        "parents": parents,
        "depends_on": depends,
        "required_for": required_for,
        "relates": relates,
        "fields": fields,
    }


# ---------------------------------------------------------------------------
# 일정·상태


def planned_dates(rec: Dict[str, Any]) -> Range:
    """계획 시작·종료. 지금은 Sprints(월 단위) 근사 — 시작일·완료일 date 필드가 생기면 여기만 바꾼다."""
    months = []
    for sprint in rec.get("sprints") or []:
        m = SPRINT_RE.match(sprint or "")
        if m and 1 <= int(m.group(2)) <= 12:
            months.append((int(m.group(1)), int(m.group(2))))
    if not months:
        return None, None
    y0, m0 = min(months)
    y1, m1 = max(months)
    return dt.date(y0, m0, 1), dt.date(y1, m1, calendar.monthrange(y1, m1)[1])


def actual_start(activities: Optional[Iterable[Dict[str, Any]]]) -> Optional[dt.date]:
    """State(`상태`)가 처음 In Progress로 바뀐 날. 이력이 없으면 None."""
    first = None
    for act in activities or []:
        if (act.get("field") or {}).get("name") not in STATE_FIELD_NAMES:
            continue
        if IN_PROGRESS not in _names(act.get("added")):
            continue
        ts = act.get("timestamp")
        if ts and (first is None or ts < first):
            first = ts
    return ms_to_date(first)


def status_bucket(rec: Dict[str, Any]) -> str:
    if rec.get("resolved"):
        return BUCKET_DONE
    if rec.get("state") in NOT_STARTED_STATES:
        return BUCKET_TODO
    return BUCKET_DOING


def is_cancelled(rec: Dict[str, Any]) -> bool:
    """resolved지만 Won't fix·Duplicate·Obsolete·Incomplete — 진척률 분모에서 뺀다."""
    return bool(rec.get("resolved")) and rec.get("state") in CANCEL_STATES


def actual_interval(row: Dict[str, Any], today: dt.date) -> Range:
    """표시용 실제 구간. 착수 후 미종료는 오늘까지, 시작 이력 없이 종료는 종료일 하루만."""
    start, end = row["act_start"], row["act_end"]
    if start and not end:
        end = max(today, start)
    if end and not start:
        start = end
    if start and end and end < start:
        start, end = end, start
    return start, end


# ---------------------------------------------------------------------------
# 트리


class Tree:
    def __init__(self, root: str) -> None:
        self.root = root
        self.issues: Dict[str, Dict[str, Any]] = {}
        self.parent: Dict[str, Optional[str]] = {root: None}  # 1차 부모 (BFS 최초 발견)
        self.children: Dict[str, List[str]] = defaultdict(list)
        self.tree_parents: Dict[str, List[str]] = defaultdict(list)  # 트리 안에서 발견된 모든 부모
        self.cycles: List[Tuple[str, str]] = []  # (부모, 이미 조상인 자식)
        self.missing: List[str] = []  # 링크는 있는데 조회 안 된 이슈
        self.order: List[str] = []  # 전위 순회
        self.depth: Dict[str, int] = {}
        self.wbs: Dict[str, str] = {}
        self.multi_parents: Dict[str, List[str]] = {}

    def is_leaf(self, iid: str) -> bool:
        return not self.children.get(iid)


def _is_ancestor(parent: Dict[str, Optional[str]], candidate: str, node: str) -> bool:
    cur: Optional[str] = node
    while cur is not None:
        if cur == candidate:
            return True
        cur = parent.get(cur)
    return False


def _sibling_key(rec: Dict[str, Any]) -> Tuple[bool, dt.date, Tuple[str, int, str]]:
    start, _ = planned_dates(rec)
    return (start is None, start or dt.date.max, id_key(rec["id"]))


def build_tree(root_id: str, fetch_issues: Callable[[List[str]], Dict[str, Dict[str, Any]]]) -> Tree:
    """BFS로 레벨마다 한 번에 조회한다. 이미 본 이슈는 다시 펼치지 않는다."""
    tree = Tree(root_id)
    frontier = [root_id]
    while frontier:
        fetched = fetch_issues(frontier)
        nxt: List[str] = []
        for iid in sorted(frontier, key=id_key):
            raw = fetched.get(iid)
            if raw is None:
                tree.missing.append(iid)
                continue
            rec = normalize_issue(raw)
            tree.issues[iid] = rec
            for cid in rec["children"]:
                if cid in tree.parent:
                    if tree.parent[cid] == iid:
                        continue  # 같은 링크 중복
                    if _is_ancestor(tree.parent, cid, iid):
                        tree.cycles.append((iid, cid))
                    elif iid not in tree.tree_parents[cid]:
                        tree.tree_parents[cid].append(iid)
                    continue
                tree.parent[cid] = iid
                tree.children[iid].append(cid)
                tree.tree_parents[cid].append(iid)
                nxt.append(cid)
        frontier = nxt

    missing = set(tree.missing)
    for pid in list(tree.children):
        kids = [c for c in tree.children[pid] if c not in missing and c in tree.issues]
        tree.children[pid] = sorted(kids, key=lambda c: _sibling_key(tree.issues[c]))

    if root_id in tree.issues:
        stack = [(root_id, "1", 1)]
        while stack:
            iid, num, depth = stack.pop()
            tree.order.append(iid)
            tree.wbs[iid] = num
            tree.depth[iid] = depth
            kids = tree.children.get(iid, [])
            for idx in range(len(kids), 0, -1):
                stack.append((kids[idx - 1], f"{num}.{idx}", depth + 1))

    for iid, rec in tree.issues.items():
        parents = set(tree.tree_parents.get(iid, [])) | set(rec["parents"])
        if len(parents) > 1:
            tree.multi_parents[iid] = sorted(parents, key=id_key)
    return tree


def compute_progress(tree: Tree) -> Dict[str, Tuple[int, int]]:
    """이슈별 (종료 leaf 수, 대상 leaf 수). 취소성 종료 leaf는 분자·분모 모두에서 뺀다."""
    result: Dict[str, Tuple[int, int]] = {}
    for iid in reversed(tree.order):
        kids = tree.children.get(iid)
        rec = tree.issues[iid]
        if not kids:
            result[iid] = (0, 0) if is_cancelled(rec) else (1 if rec["resolved"] else 0, 1)
        else:
            result[iid] = (sum(result[c][0] for c in kids), sum(result[c][1] for c in kids))
    return result


def expanded_by_default(tree: Tree, iid: str) -> bool:
    """Feature 수준까지 펼친다 — 루트는 항상 펼치고, 그 밖에 Task 하위가 하나라도 있으면 접는다."""
    kids = tree.children.get(iid) or []
    if not kids:
        return False
    if iid == tree.root:
        return True
    return all(tree.issues[c]["type"] != "Task" for c in kids)


def outline_plan(tree: Tree, collapse: bool) -> Dict[str, Dict[str, Any]]:
    """엑셀 행 그룹: outline_level = 레벨-1, collapse면 기본 접힘 그룹 아래 행을 숨긴다."""
    plan: Dict[str, Dict[str, Any]] = {}
    for iid in tree.order:
        parent = tree.parent.get(iid)
        hidden = False
        if collapse and parent is not None:
            hidden = plan[parent]["hidden"] or not expanded_by_default(tree, parent)
        plan[iid] = {
            "level": min(tree.depth[iid] - 1, 7),
            "hidden": hidden,
            "collapsed": collapse and not tree.is_leaf(iid) and not expanded_by_default(tree, iid),
        }
    return plan


# ---------------------------------------------------------------------------
# 행·rollup·점검 (엑셀·HTML 공용)


def build_rows(tree: Tree, activities: Dict[str, List[Dict[str, Any]]]) -> List[Dict[str, Any]]:
    progress = compute_progress(tree)
    rows = []
    for iid in tree.order:
        rec = tree.issues[iid]
        plan_start, plan_end = planned_dates(rec)
        done, total = progress[iid]
        rows.append({
            "id": iid,
            "wbs": tree.wbs[iid],
            "level": tree.depth[iid],
            "summary": rec["summary"],
            "type": rec["type"],
            "project": rec["project"],
            "assignee": rec["assignee"],
            "state": rec["state"],
            "bucket": status_bucket(rec),
            "cancelled": is_cancelled(rec),
            "sprints": rec["sprints"],
            "estimate": rec["estimate"],
            "spent": rec["spent"],
            "sp": rec["sp"],
            "plan_start": plan_start,
            "plan_end": plan_end,
            "act_start": actual_start(activities.get(iid)),
            "act_end": rec["resolved"],
            "depends_on": rec["depends_on"],
            "leaf_done": done,
            "leaf_total": total,
            "progress": done / total if total else None,
            "is_leaf": tree.is_leaf(iid),
        })
    return rows


def _union(a: Optional[Range], b: Optional[Range]) -> Optional[Range]:
    if not b or b[0] is None or b[1] is None:
        return a
    if not a:
        return (b[0], b[1])
    return (min(a[0], b[0]), max(a[1], b[1]))


def _wider(rollup: Optional[Range], own: Range) -> bool:
    if not rollup:
        return False
    if own[0] is None or own[1] is None:
        return True
    return rollup[0] < own[0] or rollup[1] > own[1]


def compute_rollups(tree: Tree, rows: List[Dict[str, Any]], today: dt.date) -> Dict[str, Dict[str, Any]]:
    """상위 행별 하위(취소 제외) 계획·실제 범위. 자기 일정이 없거나 하위가 더 넓으면 applied."""
    by_id = {r["id"]: r for r in rows}
    result: Dict[str, Dict[str, Any]] = {}
    for iid in reversed(tree.order):
        kids = tree.children.get(iid)
        if not kids:
            continue
        plan: Optional[Range] = None
        act: Optional[Range] = None
        for cid in kids:
            child = by_id[cid]
            if not child["cancelled"]:
                plan = _union(plan, (child["plan_start"], child["plan_end"]))
                act = _union(act, actual_interval(child, today))
            sub = result.get(cid)
            if sub:
                plan = _union(plan, sub["plan"])
                act = _union(act, sub["act"])
        row = by_id[iid]
        result[iid] = {
            "plan": plan,
            "act": act,
            "plan_applied": _wider(plan, (row["plan_start"], row["plan_end"])),
            "act_applied": _wider(act, actual_interval(row, today)),
        }
    return result


def rollup_applied_count(rollups: Dict[str, Dict[str, Any]]) -> int:
    return sum(1 for r in rollups.values() if r["plan_applied"] or r["act_applied"])


def fmt_md(d: dt.date) -> str:
    return f"{d.month}/{d.day}"


def range_label(start: Optional[dt.date], end: Optional[dt.date], ongoing: bool = False) -> str:
    """'4/1–5/31'. 진행 중이면 '4/13–진행', 시작 이력 없이 종료면 '?–5/7', 없으면 ''."""
    if ongoing and start:
        return f"{fmt_md(start)}–진행"
    if start and end:
        return f"{fmt_md(start)}–{fmt_md(end)}"
    if end:
        return f"?–{fmt_md(end)}"
    if start:
        return f"{fmt_md(start)}–?"
    return ""


def late_days(row: Dict[str, Any], today: dt.date) -> Optional[int]:
    """실제 종료(진행 중이면 기준일)가 계획 종료를 넘긴 일수. 넘기지 않았으면 None."""
    end = row["act_end"] or (today if row["act_start"] else None)
    if row["plan_end"] and end and end > row["plan_end"]:
        return (end - row["plan_end"]).days
    return None


def link_start(row: Dict[str, Any]) -> Optional[dt.date]:
    return row["act_start"] or row["plan_start"]


def link_end(row: Dict[str, Any]) -> Optional[dt.date]:
    return row["act_end"] or row["plan_end"]


def is_violation(pred: Dict[str, Any], succ: Dict[str, Any]) -> bool:
    """후행 시작(실제 우선, 없으면 계획)이 선행 종료(실제 우선, 없으면 계획)보다 앞서면 위반."""
    start, end = link_start(succ), link_end(pred)
    return bool(start and end and start < end)


def collect_links(tree: Tree) -> List[Dict[str, Any]]:
    """트리 이슈의 Depend(선행 src → 후행 dst)·Relates 링크. in_tree = 양 끝이 모두 트리 안."""
    links: List[Dict[str, Any]] = []
    seen = set()

    def add(kind: str, src: str, dst: str) -> None:
        if src == dst or (kind, src, dst) in seen:
            return
        seen.add((kind, src, dst))
        links.append({"kind": kind, "src": src, "dst": dst, "in_tree": src in tree.wbs and dst in tree.wbs})

    for iid in tree.order:
        rec = tree.issues[iid]
        for pred, _ in rec["depends_on"]:
            add(DEPEND, pred, iid)
        for succ, _ in rec["required_for"]:
            add(DEPEND, iid, succ)
        for other, _ in rec["relates"]:
            a, b = sorted((iid, other), key=id_key)
            add(RELATES, a, b)
    return links


def link_counts(links: List[Dict[str, Any]]) -> Dict[str, Dict[str, int]]:
    out: Dict[str, Dict[str, int]] = {DEPEND: {"in_tree": 0, "external": 0}, RELATES: {"in_tree": 0, "external": 0}}
    for link in links:
        out[link["kind"]]["in_tree" if link["in_tree"] else "external"] += 1
    return out


def dependency_violations(links: List[Dict[str, Any]], rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    by_id = {r["id"]: r for r in rows}
    return [link for link in links if link["kind"] == DEPEND and link["in_tree"]
            and is_violation(by_id[link["src"]], by_id[link["dst"]])]


def title_prefix(summary: str) -> Optional[str]:
    m = PREFIX_RE.match(summary or "")
    return f"[{m.group(1).strip()}]" if m else None


def compute_checks(tree: Tree, rows: List[Dict[str, Any]], today: dt.date) -> List[Dict[str, Any]]:
    """점검 항목: key, title, rule, note, items[(티켓, 비고)]."""
    checks: List[Dict[str, Any]] = []

    def add(key: str, title: str, rule: str, items: List[Tuple[str, str]], note: str = "") -> None:
        checks.append({"key": key, "title": title, "rule": rule, "note": note, "items": items})

    root_kids = tree.children.get(tree.root, [])
    flat = [c for c in root_kids if tree.is_leaf(c)]
    add("root_flat", "중간 계층 부재 (루트 직속 leaf)",
        "루트 바로 아래에 하위 없이 매달린 이슈 — Epic→Feature→Task 중간 계층이 없다는 신호",
        [(c, tree.issues[c]["type"] or "") for c in flat],
        f"루트 직속 자식 {len(root_kids)}건 중 leaf {len(flat)}건, 하위 보유 {len(root_kids) - len(flat)}건")

    prefixes = {r["id"]: title_prefix(r["summary"]) for r in rows}
    counts = Counter(p for p in prefixes.values() if p)
    major = counts.most_common(1)[0][0] if counts else None
    no_prefix = sum(1 for p in prefixes.values() if not p)
    variants = [(r["id"], prefixes[r["id"]] or "(접두어 없음)") for r in rows if prefixes[r["id"]] != major]
    note = " / ".join(f"{p} {n}" for p, n in counts.most_common())
    if no_prefix:
        note += f" / 접두어 없음 {no_prefix}"
    add("prefix_variant", "제목 접두어 변형",
        f"제목 맨 앞 [..] 표기가 다수 표기 {major or '(없음)'}와 다름", variants, note)

    add("no_sprints", "Sprints 없음 (계획 일정 공란)",
        "Sprints가 비었거나 YYYY.MM 형식이 아니라 계획 일정을 낼 수 없음",
        [(r["id"], r["bucket"]) for r in rows if r["plan_start"] is None])

    add("done_without_progress", "In Progress 이력 없이 종료",
        f"resolved인데 State가 In Progress였던 적 없음 (취소성 종료 {', '.join(sorted(CANCEL_STATES))} 제외) — 실제 시작 공란",
        [(r["id"], r["state"] or "") for r in rows
         if r["bucket"] == BUCKET_DONE and r["act_start"] is None and not r["cancelled"]])

    add("doing_without_progress", "진행 상태인데 In Progress 이력 없음",
        "미해결·착수 상태인데 State가 In Progress였던 적 없음 — 실제 시작 공란",
        [(r["id"], r["state"] or "") for r in rows if r["bucket"] == BUCKET_DOING and r["act_start"] is None])

    blocked = []
    for r in rows:
        if r["bucket"] != BUCKET_DOING:
            continue
        open_deps = [d for d, resolved in r["depends_on"] if not resolved]
        if open_deps:
            blocked.append((r["id"], "선행 미종료: " + ", ".join(open_deps)))
    add("blocked_in_progress", "선행 미종료인데 진행중",
        "depends on 대상이 미해결인데 이 이슈는 진행 상태", blocked)

    by_id = {r["id"]: r for r in rows}
    add("dependency_violation", "선후 위반",
        "Depend 후행의 시작(실제 우선, 없으면 계획)이 선행의 종료(실제 우선, 없으면 계획)보다 앞섬 — 트리 안 링크만",
        [(link["dst"], f"선행 {link['src']} 종료 {link_end(by_id[link['src']])} > 후행 시작 "
                       f"{link_start(by_id[link['dst']])}")
         for link in dependency_violations(collect_links(tree), rows)])

    add("multi_parent", "다중 부모",
        "Subtask 부모가 둘 이상 — WBS에는 첫 부모 아래 한 번만 표시",
        [(iid, "부모: " + ", ".join(ps)) for iid, ps in sorted(tree.multi_parents.items(), key=lambda x: id_key(x[0]))])

    add("overdue", "미해결인데 Sprints가 지난 달",
        f"미해결이고 계획 종료(마지막 Sprint 말일)가 기준일 {today.isoformat()} 이전",
        [(r["id"], f"{r['state'] or ''} / 계획 종료 {r['plan_end'].isoformat()}") for r in rows
         if r["bucket"] != BUCKET_DONE and r["plan_end"] and r["plan_end"] < today])

    add("cancelled", "취소성 종료",
        f"resolved지만 State가 {', '.join(sorted(CANCEL_STATES))} — 진척률 분모에서 제외",
        [(r["id"], r["state"] or "") for r in rows if r["cancelled"]])

    add("cycle", "순환 링크", "Subtask 링크가 조상으로 되돌아감 — 한 번만 펼침",
        [(child, f"{parent} → {child}") for parent, child in tree.cycles])

    add("missing", "조회 불가", "Subtask 링크는 있는데 조회되지 않음 (권한·삭제)",
        [(iid, "") for iid in tree.missing])
    return checks


def team_summary(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    groups: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for r in rows:
        groups[r["project"]].append(r)
    order = sorted(groups, key=lambda p: (-len(groups[p]), p))
    out = []
    for label, items in [(p, groups[p]) for p in order] + [("전체", rows)]:
        leaves = [r for r in items if r["is_leaf"] and not r["cancelled"]]
        leaf_done = sum(1 for r in leaves if r["bucket"] == BUCKET_DONE)
        out.append({
            "project": label,
            "count": len(items),
            "done": sum(1 for r in items if r["bucket"] == BUCKET_DONE),
            "cancelled": sum(1 for r in items if r["cancelled"]),
            "doing": sum(1 for r in items if r["bucket"] == BUCKET_DOING),
            "todo": sum(1 for r in items if r["bucket"] == BUCKET_TODO),
            "leaves": len(leaves),
            "leaf_done": leaf_done,
            "progress": leaf_done / len(leaves) if leaves else None,
            "no_plan": sum(1 for r in items if r["plan_start"] is None),
            "no_actual": sum(1 for r in items if r["bucket"] != BUCKET_TODO and r["act_start"] is None
                             and not r["cancelled"]),
        })
    return out


# ---------------------------------------------------------------------------
# 타임라인


def monday(d: dt.date) -> dt.date:
    return d - dt.timedelta(days=d.weekday())


def week_starts(start: dt.date, end: dt.date) -> List[dt.date]:
    weeks = []
    cur = monday(start)
    while cur <= end:
        weeks.append(cur)
        cur += dt.timedelta(days=7)
    return weeks


def is_weekend(d: dt.date) -> bool:
    return d.weekday() >= 5


def timeline_range(rows: List[Dict[str, Any]], rollups: Dict[str, Dict[str, Any]], today: dt.date) -> Range:
    """간트 날짜 범위 = 계획·실제·rollup·기준일의 최소~최대."""
    dates = [today]
    for r in rows:
        dates.extend(d for d in (r["plan_start"], r["plan_end"]) if d)
        dates.extend(d for d in actual_interval(r, today) if d)
    for ru in rollups.values():
        for rng in (ru["plan"], ru["act"]):
            if rng:
                dates.extend(rng)
    return min(dates), max(dates)


def timeline_columns(start: dt.date, end: dt.date, unit: str) -> List[Tuple[dt.date, dt.date]]:
    """간트 열 목록 (열 시작일, 열 종료일). day는 하루 한 열, week는 월요일 시작 7일 한 열."""
    if unit == "week":
        return [(w, w + dt.timedelta(days=6)) for w in week_starts(start, end)]
    return [(start + dt.timedelta(days=i), start + dt.timedelta(days=i)) for i in range((end - start).days + 1)]


def _overlaps(rng: Optional[Range], c0: dt.date, c1: dt.date) -> bool:
    return bool(rng and rng[0] and rng[1] and rng[0] <= c1 and rng[1] >= c0)


def cell_kind(row: Dict[str, Any], actual: Range, rollup: Optional[Dict[str, Any]],
              c0: dt.date, c1: dt.date) -> Optional[str]:
    """간트 한 칸의 종류. 우선순위: 실제 > 하위 집계 실제 > 계획 > 하위 집계 계획 (실제가 계획보다 먼저)."""
    planned = _overlaps((row["plan_start"], row["plan_end"]), c0, c1)
    act = _overlaps(actual, c0, c1)
    if row["cancelled"]:
        return "cancel_act" if act else ("cancel_plan" if planned else None)
    if act:
        return "act_in" if planned else "act_out"
    if rollup and rollup["act_applied"] and _overlaps(rollup["act"], c0, c1):
        return "ru_act"
    if planned:
        return "plan"
    if rollup and rollup["plan_applied"] and _overlaps(rollup["plan"], c0, c1):
        return "ru_plan"
    return None


# ---------------------------------------------------------------------------
# 엑셀


def team_label(project: str) -> str:
    return f"{project} {TEAM_NAMES[project]}" if project in TEAM_NAMES else project


def issue_url(base_url: str, iid: str) -> str:
    return f"{base_url}/issue/{iid}"


def type_color_key(type_name: Optional[str]) -> str:
    """유형 이름 → TYPE_COLORS 키. 표에 없는 유형·빈 값은 중립 회색(fallback)."""
    key = (type_name or "").strip().lower()
    return key if key in TYPE_COLORS else TYPE_COLOR_FALLBACK


def write_workbook(path: Path, tree: Tree, rows: List[Dict[str, Any]], checks: List[Dict[str, Any]],
                   rollups: Dict[str, Dict[str, Any]], meta: List[Tuple[str, Any]], base_url: str,
                   today: dt.date, unit: str = "day") -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import Outline

    def solid(color: str) -> PatternFill:
        return PatternFill("solid", fgColor=color)

    header_fill = solid("404040")
    header_font = Font(bold=True, color="FFFFFF")
    link_font = Font(color="0563C1", underline="single")
    link_font_l1 = Font(color="FFFFFF", underline="single", bold=True)
    red_font = Font(color="C00000")
    bold = Font(bold=True)
    level_fill = {1: solid("1F3864"), 2: solid("D9E2F3")}
    level_font = {1: Font(bold=True, color="FFFFFF"), 2: Font(bold=True)}
    type_fill = {k: solid(c["bg"]) for k, c in TYPE_COLORS.items()}
    kind_fill = {
        "plan": solid("BDD7EE"),
        "act_in": solid("2F5597"),
        "act_out": solid("ED7D31"),
        "ru_plan": PatternFill("lightUp", fgColor="5B9BD5", bgColor="FFFFFF"),
        "ru_act": PatternFill("darkUp", fgColor="2F5597", bgColor="BDD7EE"),
        "cancel_plan": solid("E7E6E6"),
        "cancel_act": solid("A6A6A6"),
    }
    weekend_fill = solid("F2F2F2")
    today_fill = solid("C00000")
    red_side = Side(style="medium", color="C00000")
    today_border = Border(left=red_side, right=red_side)
    date_fmt = "yyyy-mm-dd"

    def header(ws, row: int, labels: List[str]) -> None:
        for col, label in enumerate(labels, 1):
            cell = ws.cell(row=row, column=col, value=label)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def link(cell, iid: str, level: int = 0) -> None:
        cell.value = iid
        cell.hyperlink = issue_url(base_url, iid)
        cell.font = link_font_l1 if level == 1 else link_font

    def widths(ws, values: List[float]) -> None:
        for idx, w in enumerate(values, 1):
            ws.column_dimensions[get_column_letter(idx)].width = w

    def level_row(ws, row_no: int, level: int, ncols: int) -> None:
        if level in level_fill:
            for col in range(1, ncols + 1):
                cell = ws.cell(row=row_no, column=col)
                cell.fill = level_fill[level]
                if not cell.hyperlink:
                    cell.font = level_font[level]

    def apply_outline(ws, first_row: int, collapse: bool) -> None:
        plan = outline_plan(tree, collapse)
        ws.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=True)
        for offset, iid in enumerate(tree.order):
            dim = ws.row_dimensions[first_row + offset]
            dim.outline_level = plan[iid]["level"]
            dim.hidden = plan[iid]["hidden"]
            dim.collapsed = plan[iid]["collapsed"]
        ws.sheet_format.outlineLevelRow = max((p["level"] for p in plan.values()), default=0)

    def progress_text(r: Dict[str, Any]) -> str:
        return f"{r['leaf_done']}/{r['leaf_total']}" if r["leaf_total"] else "-"

    wb = Workbook()

    # 1. WBS — 전체 펼침, 레벨대로 그룹만 건다
    ws = wb.active
    ws.title = "WBS"
    labels = WBS_LABELS
    header(ws, 1, labels)
    for i, r in enumerate(rows, 2):
        deps = ", ".join(d if resolved else f"{d}(미종료)" for d, resolved in r["depends_on"])
        values = [r["wbs"], r["level"], None, r["summary"], r["type"], team_label(r["project"]), r["assignee"],
                  r["state"], BUCKET_CANCEL if r["cancelled"] else r["bucket"], ", ".join(r["sprints"]),
                  r["estimate"], r["spent"], r["sp"], r["plan_start"], r["plan_end"], r["act_start"],
                  r["act_end"], deps, r["progress"], progress_text(r)]
        for col, value in enumerate(values, 1):
            ws.cell(row=i, column=col, value=value)
        link(ws.cell(row=i, column=3), r["id"], r["level"])
        level_row(ws, i, r["level"], len(labels))
        type_cell, tc = ws.cell(row=i, column=5), type_color_key(r["type"])  # 레벨 채움 뒤에 칠해야 유형 색이 남는다
        type_cell.fill = type_fill[tc]
        type_cell.font = Font(bold=r["level"] in level_font, color=TYPE_COLORS[tc]["ink"])
        type_cell.alignment = Alignment(horizontal="center")
        ws.cell(row=i, column=4).alignment = Alignment(indent=max(r["level"] - 1, 0))
        for col in (14, 15, 16, 17):
            ws.cell(row=i, column=col).number_format = date_fmt
        if r["bucket"] != BUCKET_DONE and r["plan_end"] and r["plan_end"] < today:
            ws.cell(row=i, column=15).font = Font(color="FFB3B3", bold=True) if r["level"] == 1 else red_font
        ws.cell(row=i, column=19).number_format = "0%"
    apply_outline(ws, 2, collapse=False)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(labels))}{len(rows) + 1}"
    widths(ws, [10, 5, 12, 60, 9, 13, 9, 11, 7, 18, 9, 12, 5, 11, 11, 11, 11, 24, 9, 9])

    # 2. 간트 — 날짜(또는 주) 열 칸 색칠, Task 그룹은 접힘
    gs = wb.create_sheet("간트")
    fixed = ["WBS", "티켓", "제목", "팀", "상태", "진척률"]
    first_col = len(fixed) + 1
    first_row = 6
    start, end = timeline_range(rows, rollups, today)
    columns = timeline_columns(start, end, unit)
    legend = [
        (1, 2, "계획", "plan", None), (1, 3, "실제 · 계획 안", "act_in", "FFFFFF"),
        (1, 4, "실제 · 계획 밖", "act_out", "FFFFFF"), (1, 5, "취소", "cancel_act", None),
        (1, 6, "오늘", None, "FFFFFF"), (2, 2, "하위 집계 계획", "ru_plan", None),
        (2, 3, "하위 집계 실제 (빗금 = 하위에서 집계)", "ru_act", "FFFFFF"), (2, 4, "주말", None, None),
    ]
    gs.cell(row=1, column=1, value="범례").font = bold
    gs.cell(row=2, column=1, value=f"단위: {'일' if unit == 'day' else '주'}").font = bold
    for row_no, col, text, kind, color in legend:
        cell = gs.cell(row=row_no, column=col, value=text)
        cell.fill = kind_fill[kind] if kind else (today_fill if text == "오늘" else weekend_fill)
        cell.font = Font(bold=True, color=color) if color else bold
        cell.alignment = Alignment(horizontal="center")
    header(gs, 3, fixed)
    for col in range(1, len(fixed) + 1):
        gs.cell(row=4, column=col).fill = header_fill
        gs.cell(row=5, column=col).fill = header_fill
        gs.merge_cells(start_row=3, start_column=col, end_row=5, end_column=col)

    def header_runs(row_no: int, keys: List[Any], labels_: List[str]) -> None:
        run_start = 0
        for idx in range(len(keys) + 1):
            if idx == len(keys) or keys[idx] != keys[run_start]:
                c_from, c_to = first_col + run_start, first_col + idx - 1
                cell = gs.cell(row=row_no, column=c_from, value=labels_[run_start])
                cell.alignment = Alignment(horizontal="left" if c_to > c_from else "center")
                if c_to > c_from:
                    gs.merge_cells(start_row=row_no, start_column=c_from, end_row=row_no, end_column=c_to)
                run_start = idx
        for idx in range(len(keys)):
            cell = gs.cell(row=row_no, column=first_col + idx)
            cell.fill = header_fill
            cell.font = header_font

    if unit == "week":
        months = [(c0 + dt.timedelta(days=3)).strftime("%Y-%m") for c0, _ in columns]  # 주의 목요일 기준
        header_runs(3, months, months)
        header_runs(4, [c0 for c0, _ in columns], [c0.strftime("%m/%d") for c0, _ in columns])
        header_runs(5, list(range(len(columns))), [f"W{c0.isocalendar()[1]}" for c0, _ in columns])
    else:
        months = [c0.strftime("%Y-%m") for c0, _ in columns]
        header_runs(3, months, months)
        mondays = [monday(c0) for c0, _ in columns]
        header_runs(4, mondays, [m.strftime("%m/%d") for m in mondays])
        header_runs(5, list(range(len(columns))), [str(c0.day) for c0, _ in columns])
    for idx, (c0, c1) in enumerate(columns):
        col = first_col + idx
        gs.column_dimensions[get_column_letter(col)].width = 2.7 if unit == "day" else 3.6
        day_cell = gs.cell(row=5, column=col)
        if c0 <= today <= c1:
            day_cell.fill = today_fill
        elif unit == "day" and is_weekend(c0):
            day_cell.font = Font(bold=True, color="BFBFBF")

    for i, r in enumerate(rows, first_row):
        gs.cell(row=i, column=1, value=r["wbs"])
        link(gs.cell(row=i, column=2), r["id"], r["level"])
        gs.cell(row=i, column=3, value=r["summary"])
        gs.cell(row=i, column=4, value=TEAM_NAMES.get(r["project"], r["project"]))
        gs.cell(row=i, column=5, value=r["state"])
        gs.cell(row=i, column=6, value=r["progress"]).number_format = "0%"
        level_row(gs, i, r["level"], len(fixed))
        gs.cell(row=i, column=3).alignment = Alignment(indent=max(r["level"] - 1, 0))
        actual = actual_interval(r, today)
        rollup = rollups.get(r["id"])
        for idx, (c0, c1) in enumerate(columns):
            kind = cell_kind(r, actual, rollup, c0, c1)
            if kind is None and not (unit == "day" and is_weekend(c0)) and not (c0 <= today <= c1):
                continue
            cell = gs.cell(row=i, column=first_col + idx)
            if kind:
                cell.fill = kind_fill[kind]
            elif unit == "day" and is_weekend(c0):
                cell.fill = weekend_fill
            if c0 <= today <= c1:
                cell.border = today_border
    apply_outline(gs, first_row, collapse=True)
    gs.freeze_panes = gs.cell(row=first_row, column=first_col)
    widths(gs, [10, 12, 48, 13, 11, 8])

    # 3. 팀별 요약
    ts = wb.create_sheet("팀별 요약")
    labels = ["프로젝트(팀)", "이슈 수", "종료", "└ 취소성 종료", "진행", "미착수", "leaf 수(취소 제외)",
              "leaf 종료", "진척률(취소 제외)", "일정 없음(Sprints)", "실제 시작 없음(착수 이후, 취소 제외)"]
    header(ts, 1, labels)
    for i, s in enumerate(team_summary(rows), 2):
        values = [team_label(s["project"]), s["count"], s["done"], s["cancelled"], s["doing"], s["todo"],
                  s["leaves"], s["leaf_done"], s["progress"], s["no_plan"], s["no_actual"]]
        for col, value in enumerate(values, 1):
            cell = ts.cell(row=i, column=col, value=value)
            if s["project"] == "전체":
                cell.font = bold
        ts.cell(row=i, column=9).number_format = "0%"
    ts.freeze_panes = "B2"
    widths(ts, [16, 8, 7, 12, 7, 8, 14, 9, 14, 16, 28])

    # 4. 점검
    cs = wb.create_sheet("점검")
    header(cs, 1, ["항목", "건수", "판정 기준", "비고"])
    for i, check in enumerate(checks, 2):
        cs.cell(row=i, column=1, value=check["title"]).font = bold
        cs.cell(row=i, column=2, value=len(check["items"]))
        cs.cell(row=i, column=3, value=check["rule"]).alignment = Alignment(wrap_text=True, vertical="top")
        cs.cell(row=i, column=4, value=check["note"]).alignment = Alignment(wrap_text=True, vertical="top")
    detail_top = len(checks) + 3
    header(cs, detail_top, ["항목", "티켓", "제목", "프로젝트(팀)", "상태", "비고"])
    row_no = detail_top + 1
    for check in checks:
        for iid, note in check["items"]:
            rec = tree.issues.get(iid, {})
            cs.cell(row=row_no, column=1, value=check["title"])
            link(cs.cell(row=row_no, column=2), iid)
            cs.cell(row=row_no, column=3, value=rec.get("summary"))
            cs.cell(row=row_no, column=4, value=team_label(rec["project"]) if rec else None)
            cs.cell(row=row_no, column=5, value=rec.get("state"))
            cs.cell(row=row_no, column=6, value=note)
            row_no += 1
    cs.auto_filter.ref = f"A{detail_top}:F{max(row_no - 1, detail_top)}"
    widths(cs, [30, 12, 60, 14, 11, 40])

    # 5. 메타
    ms = wb.create_sheet("메타")
    header(ms, 1, ["항목", "값"])
    for i, (key, value) in enumerate(meta, 2):
        ms.cell(row=i, column=1, value=key).font = bold
        ms.cell(row=i, column=2, value=value).alignment = Alignment(wrap_text=True, vertical="top")
    widths(ms, [22, 100])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


# ---------------------------------------------------------------------------
# HTML 간트 (같은 행·rollup·점검 결과를 인라인 JSON으로)


def _iso(d: Optional[dt.date]) -> Optional[str]:
    return d.isoformat() if d else None


def build_html_payload(tree: Tree, rows: List[Dict[str, Any]], checks: List[Dict[str, Any]],
                       rollups: Dict[str, Dict[str, Any]], base_url: str, today: dt.date,
                       meta: Dict[str, Any]) -> Dict[str, Any]:
    total = team_summary(rows)[-1]
    start, end = timeline_range(rows, rollups, today)
    by_key = {c["key"]: c for c in checks}
    links = collect_links(tree)
    by_id = {r["id"]: r for r in rows}
    ext: Dict[str, List[List[str]]] = defaultdict(list)
    for link in links:
        if link["in_tree"]:
            continue
        kind = "d" if link["kind"] == DEPEND else "r"
        if link["src"] in tree.wbs:
            ext[link["src"]].append(["→" if kind == "d" else "↔", link["dst"], kind])
        if link["dst"] in tree.wbs:
            ext[link["dst"]].append(["←" if kind == "d" else "↔", link["src"], kind])
    out_rows = []
    for r in rows:
        a0, a1 = actual_interval(r, today)
        ru = rollups.get(r["id"]) or {}
        ongoing = bool(r["act_start"] and not r["act_end"])
        own_plan = range_label(r["plan_start"], r["plan_end"])
        own_act = range_label(r["act_start"], r["act_end"], ongoing)
        out_rows.append({
            "id": r["id"], "wbs": r["wbs"], "lv": r["level"], "par": tree.parent.get(r["id"]),
            "title": r["summary"], "type": r["type"], "tc": type_color_key(r["type"]), "proj": r["project"],
            "team": TEAM_NAMES.get(r["project"], r["project"]), "asg": r["assignee"], "state": r["state"],
            "b": BUCKET_CANCEL if r["cancelled"] else r["bucket"],
            "spr": ", ".join(r["sprints"]), "est": r["estimate"], "spent": r["spent"], "sp": r["sp"],
            "ps": _iso(r["plan_start"]), "pe": _iso(r["plan_end"]), "as": _iso(a0), "ae": _iso(a1),
            "s0": _iso(r["act_start"]), "s1": _iso(r["act_end"]), "ongoing": bool(r["act_start"] and not r["act_end"]),
            "rp": [_iso(x) for x in ru["plan"]] if ru.get("plan_applied") else None,
            "ra": [_iso(x) for x in ru["act"]] if ru.get("act_applied") else None,
            "deps": [[d, ok] for d, ok in r["depends_on"]], "leaf": r["is_leaf"],
            "prog": r["progress"], "done": r["leaf_done"], "tot": r["leaf_total"],
            "exp": expanded_by_default(tree, r["id"]),
            "pl": own_plan or (range_label(*ru["plan"]) if ru.get("plan_applied") else ""),
            "plr": not own_plan and bool(ru.get("plan_applied")),
            "al": own_act or (range_label(*ru["act"]) if ru.get("act_applied") else ""),
            "alr": not own_act and bool(ru.get("act_applied")),
            "late": late_days(r, today),
            "ext": ext.get(r["id"], []),
        })
    return {
        "meta": dict(meta, issues=len(rows), rollups=rollup_applied_count(rollups)),
        "base": base_url,
        "today": today.isoformat(),
        "range": [_iso(start), _iso(end)],
        "kpi": {
            "total": total["count"], "done": total["done"] - total["cancelled"], "doing": total["doing"],
            "todo": total["todo"], "cancelled": total["cancelled"], "progress": total["progress"],
            "leaf_done": total["leaf_done"], "leaves": total["leaves"],
        },
        "teams": [{"proj": p, "team": TEAM_NAMES.get(p, p), "count": n}
                  for p, n in sorted(Counter(r["project"] for r in rows).items(), key=lambda x: -x[1])],
        "summary": [dict(s, label=team_label(s["project"])) for s in team_summary(rows)],
        "checks": [{"key": c["key"], "title": c["title"], "rule": c["rule"], "note": c["note"],
                    "items": [[i, note] for i, note in c["items"]]} for c in checks],
        "pinned": [k for k in CHIP_KEYS if k in by_key],
        "type_colors": {k: {name: "#" + hex_ for name, hex_ in c.items()} for k, c in TYPE_COLORS.items()},
        "rows": out_rows,
        "links": [{"k": "d" if link["kind"] == DEPEND else "r", "f": link["src"], "t": link["dst"],
                   "v": link["kind"] == DEPEND and is_violation(by_id[link["src"]], by_id[link["dst"]])}
                  for link in links if link["in_tree"]],
        "link_counts": link_counts(links),
    }


def render_html(payload: Dict[str, Any], title: str) -> str:
    template = HTML_TEMPLATE.read_text(encoding="utf-8")
    data = json.dumps(payload, ensure_ascii=False).replace("<", "\\u003c")
    return template.replace("__TITLE__", html.escape(title)).replace("/*__DATA__*/", data)


# ---------------------------------------------------------------------------
# YouTrack (GET 전용)


class YouTrackError(RuntimeError):
    pass


class YouTrackClient:
    def __init__(self, base_url: str, token: str, timeout: int = 60) -> None:
        self.base_url = base_url.rstrip("/")
        self._token = token
        self.timeout = timeout
        self.calls = 0
        self._lock = threading.Lock()

    def get(self, path: str, params: Dict[str, str]) -> Any:
        url = f"{self.base_url}{path}?{urllib.parse.urlencode(params)}"
        req = urllib.request.Request(url, method="GET", headers={
            "Authorization": f"Bearer {self._token}",
            "Accept": "application/json",
        })
        for attempt in range(RETRIES + 1):  # GET은 멱등이라 일시 오류(타임아웃·429·5xx)는 재시도한다
            with self._lock:
                self.calls += 1
            try:
                with urllib.request.urlopen(req, timeout=self.timeout) as resp:
                    return json.load(resp)
            except urllib.error.HTTPError as exc:
                if exc.code not in (429, 500, 502, 503, 504) or attempt == RETRIES:
                    raise YouTrackError(f"GET {path} → HTTP {exc.code}") from None
            except (urllib.error.URLError, OSError) as exc:
                if attempt == RETRIES:
                    raise YouTrackError(f"GET {path} → {getattr(exc, 'reason', exc)}") from None
            time.sleep(1 + attempt)
        raise YouTrackError(f"GET {path} → 재시도 초과")

    def fetch_issues(self, ids: List[str]) -> Dict[str, Dict[str, Any]]:
        out: Dict[str, Dict[str, Any]] = {}
        for i in range(0, len(ids), ISSUE_BATCH):
            chunk = ids[i:i + ISSUE_BATCH]
            try:
                items = self.get("/api/issues", {
                    "query": "issue id: " + ", ".join(chunk),
                    "fields": ISSUE_FIELDS,
                    "$top": str(len(chunk) + 10),
                })
            except YouTrackError:
                items = []  # 질의 실패(존재하지 않는 ID 등)면 건별로 재시도
                for iid in chunk:
                    try:
                        items.append(self.get(f"/api/issues/{urllib.parse.quote(iid)}", {"fields": ISSUE_FIELDS}))
                    except YouTrackError:
                        pass
            for item in items:
                if item.get("idReadable") in chunk:
                    out[item["idReadable"]] = item
        return out

    def fetch_activities(self, iid: str) -> List[Dict[str, Any]]:
        items = self.get(f"/api/issues/{urllib.parse.quote(iid)}/activities", {
            "categories": "CustomFieldCategory",
            "fields": ACTIVITY_FIELDS,
            "$top": str(ACTIVITY_TOP),
        })
        if len(items) >= ACTIVITY_TOP:
            print(f"경고: {iid} 활동이 {ACTIVITY_TOP}건 이상 — 일부 누락 가능", file=sys.stderr)
        return items


def resolve_token() -> str:
    token = os.environ.get("YOUTRACK_TOKEN")
    if token:
        return token
    proc = subprocess.run([sys.executable, str(REPO / "tools" / "cred.py"), "get", TOKEN_CRED_NAME],
                          capture_output=True, text=True)
    token = proc.stdout.strip()
    if proc.returncode != 0 or not token:
        raise YouTrackError(f"YouTrack 토큰 없음 — python3 tools/cred.py set {TOKEN_CRED_NAME}")
    return token


def ensure_outside_repo(out: Path) -> Path:
    """산출물은 업무 데이터라 repo 안에 쓰지 않는다."""
    resolved = out.expanduser().resolve()
    if resolved == REPO or REPO in resolved.parents:
        raise ValueError(f"산출물 경로는 repo 밖이어야 한다: {resolved}")
    return resolved


def depth_distribution(tree: Tree) -> Dict[int, int]:
    return dict(sorted(Counter(tree.depth.values()).items()))


def parse_args(argv: Optional[List[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="YouTrack 루트 이슈 하위 트리 → WBS·간트 엑셀 (+HTML 간트)")
    parser.add_argument("root", help="루트 이슈 ID (예: DEV2-5237)")
    parser.add_argument("--out", required=True, help="출력 .xlsx 경로 (repo 밖)")
    parser.add_argument("--html", help="HTML 간트 출력 경로 (repo 밖, 선택)")
    parser.add_argument("--html-title", help="HTML <title> (기본: '{루트} 간트')")
    parser.add_argument("--gantt-unit", choices=("day", "week"), default="day", help="엑셀 간트 열 단위 (기본 day)")
    parser.add_argument("--today", help="기준일 YYYY-MM-DD (기본: 오늘, KST)")
    parser.add_argument("--workers", type=int, default=8, help="활동 조회 동시 요청 수 (기본 8)")
    return parser.parse_args(argv)


def main(argv: Optional[List[str]] = None) -> int:
    args = parse_args(argv)
    try:
        out = ensure_outside_repo(Path(args.out))
        html_out = ensure_outside_repo(Path(args.html)) if args.html else None
    except ValueError as exc:
        print(f"오류: {exc}", file=sys.stderr)
        return 2
    today = dt.date.fromisoformat(args.today) if args.today else dt.datetime.now(KST).date()
    base_url = os.environ.get("YOUTRACK_BASE_URL", DEFAULT_BASE_URL).rstrip("/")
    started = time.monotonic()
    try:
        client = YouTrackClient(base_url, resolve_token())
        tree = build_tree(args.root, client.fetch_issues)
        if args.root not in tree.issues:
            print(f"오류: 루트 이슈 {args.root}를 조회할 수 없다", file=sys.stderr)
            return 1
        with ThreadPoolExecutor(max_workers=max(1, args.workers)) as pool:
            activities = dict(zip(tree.order, pool.map(client.fetch_activities, tree.order)))
    except YouTrackError as exc:
        print(f"오류: {exc}", file=sys.stderr)
        return 1

    rows = build_rows(tree, activities)
    checks = compute_checks(tree, rows, today)
    rollups = compute_rollups(tree, rows, today)
    counts = link_counts(collect_links(tree))
    elapsed = time.monotonic() - started
    depths = depth_distribution(tree)
    total = team_summary(rows)[-1]
    g_start, g_end = timeline_range(rows, rollups, today)
    extracted = dt.datetime.now(KST).strftime("%Y-%m-%d %H:%M:%S KST")
    meta = [
        ("루트 이슈", args.root),
        ("추출 시각", extracted),
        ("기준일", today.isoformat()),
        ("이슈 수", len(rows)),
        ("깊이 분포(레벨: 건수)", ", ".join(f"{k}: {v}" for k, v in depths.items())),
        ("API 호출 수", client.calls),
        ("소요 시간(초)", round(elapsed, 1)),
        ("YouTrack", base_url),
        ("계획 일정", "Sprints(월 단위) 최소월 1일 ~ 최대월 말일 근사 — 시작일·완료일 필드 없음"),
        ("실제 일정", "시작 = State가 처음 In Progress가 된 날 / 종료 = resolved 날짜"),
        ("구분", "종료 = resolved 있음(그중 " + ", ".join(sorted(CANCEL_STATES)) + "는 취소) / 미착수 = "
         + ", ".join(sorted(NOT_STARTED_STATES)) + " / 진행 = 그 외 미해결"),
        ("진척률", "하위 leaf 중 종료 비율 — 취소성 종료 leaf는 분자·분모에서 뺀다"),
        ("WBS 형제 정렬", "계획 시작 순, 같으면 티켓 ID 순. 다중 부모는 첫 부모 아래 한 번만"),
        ("간트", f"단위 {args.gantt_unit}, 범위 {g_start}~{g_end}. 행 그룹은 Feature까지 펼침·Task 접힘. "
                 "칸 우선순위 실제 > 하위 집계 실제 > 계획 > 하위 집계 계획. 진행중 실제는 기준일까지"),
        ("하위 집계(rollup)", f"자기 Sprints가 없거나 하위(취소 제외) 범위가 더 넓은 상위 행 "
                              f"{rollup_applied_count(rollups)}건 — 빗금으로 구분, 엑셀·HTML 공용 계산"),
        ("연결(링크)", f"Depend 트리 안 {counts[DEPEND]['in_tree']} / 밖 {counts[DEPEND]['external']}, "
                       f"Relates 트리 안 {counts[RELATES]['in_tree']} / 밖 {counts[RELATES]['external']}"),
    ]
    write_workbook(out, tree, rows, checks, rollups, meta, base_url, today, args.gantt_unit)
    if html_out:
        payload = build_html_payload(tree, rows, checks, rollups, base_url, today, {
            "root": args.root, "root_title": tree.issues[args.root]["summary"], "extracted": extracted,
            "api_calls": client.calls,
        })
        html_out.parent.mkdir(parents=True, exist_ok=True)
        html_out.write_text(render_html(payload, args.html_title or f"{args.root} 간트"), encoding="utf-8")

    print(f"루트 {args.root}: 이슈 {len(rows)}건, 깊이 {depths}, API 호출 {client.calls}회, {elapsed:.1f}초")
    print(f"프로젝트별: {dict(Counter(r['project'] for r in rows))}")
    progress = total["progress"]
    print(f"진척률(leaf, 취소 제외): {total['leaf_done']}/{total['leaves']} = "
          f"{'-' if progress is None else format(progress, '.1%')}")
    print(f"간트: {args.gantt_unit} 단위 {g_start}~{g_end}, rollup 적용 {rollup_applied_count(rollups)}행")
    print(f"링크: Depend 트리 안 {counts[DEPEND]['in_tree']} / 밖 {counts[DEPEND]['external']}, "
          f"Relates 트리 안 {counts[RELATES]['in_tree']} / 밖 {counts[RELATES]['external']}")
    for check in checks:
        print(f"  점검 {check['title']}: {len(check['items'])}")
    print(f"저장: {out}" + (f", {html_out}" if html_out else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main())
