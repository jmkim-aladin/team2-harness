from __future__ import annotations

import datetime as dt
import importlib.util
import json
import re
import tempfile
import unittest
from pathlib import Path


MODULE_PATH = Path(__file__).resolve().parents[1] / "tools" / "export_wbs.py"
spec = importlib.util.spec_from_file_location("export_wbs", MODULE_PATH)
wbs = importlib.util.module_from_spec(spec)
assert spec and spec.loader
spec.loader.exec_module(wbs)

TODAY = dt.date(2026, 10, 7)
DONE_MS = 1789009200000  # 2026-09-10 12:00 KST
SEP1_MS = 1788188400000  # 2026-09-01 00:00 KST
BASE = "https://example.youtrack.cloud"


def raw_issue(iid, *, children=(), parents=(), depends=(), required_for=(), relates=(), type_="Feature",
              state="Open", sprints=(), resolved=None, summary=None, project=None):
    """YouTrack /api/issues 응답 모양의 픽스처."""
    links = [
        {"direction": "OUTWARD", "linkType": {"name": "Subtask"},
         "issues": [{"idReadable": c} for c in children]},
        {"direction": "INWARD", "linkType": {"name": "Subtask"},
         "issues": [{"idReadable": p} for p in parents]},
        {"direction": "INWARD", "linkType": {"name": "Depend"},
         "issues": [{"idReadable": d, "resolved": DONE_MS if ok else None} for d, ok in depends]},
        {"direction": "OUTWARD", "linkType": {"name": "Depend"},
         "issues": [{"idReadable": d, "resolved": None} for d in required_for]},
        {"direction": "BOTH", "linkType": {"name": "Relates"},
         "issues": [{"idReadable": d, "resolved": None} for d in relates]},
    ]
    return {
        "idReadable": iid,
        "summary": summary if summary is not None else f"[스토어프론트] {iid}",
        "resolved": resolved,
        "project": {"shortName": project or iid.rpartition("-")[0]},
        "customFields": [
            {"name": "Type", "value": {"name": type_}},
            {"name": "State", "value": {"name": state}},
            {"name": "Assignee", "value": {"login": "jm", "fullName": "김정민"}},
            {"name": "Sprints", "value": [{"name": s} for s in sprints]},
            {"name": "Story points", "value": 3},
            {"name": "예측", "value": {"presentation": "3일", "minutes": 1440}},
            {"name": "소요 시간", "value": None},
        ],
        "links": links,
    }


class FakeFetcher:
    def __init__(self, issues):
        self.issues = {i["idReadable"]: i for i in issues}
        self.calls = []

    def __call__(self, ids):
        self.calls.append(list(ids))
        return {i: self.issues[i] for i in ids if i in self.issues}


def state_act(ts, field, added):
    return {"timestamp": ts, "field": {"name": field}, "added": [{"name": added}], "removed": []}


def sample_tree():
    """루트 → 직속 leaf 2(하나는 취소) + Feature(Task 2). 점검·rollup 항목을 하나씩 건드리도록 구성."""
    fetch = FakeFetcher([
        raw_issue("DEV2-1", children=["DEV2-2", "DEV2-3", "DEV2-4"], type_="Epic", state="In Progress",
                  sprints=["2026.08", "2026.10"], summary="[Store Front] 루트"),
        raw_issue("DEV2-2", parents=["DEV2-1"], state="Fixed", resolved=DONE_MS, sprints=["2026.08"]),
        raw_issue("DEV2-3", parents=["DEV2-1"], state="Won't fix", resolved=DONE_MS,
                  summary="[스토어 프론트] 취소"),
        raw_issue("DEV2-4", parents=["DEV2-1"], children=["DEV2-5", "DEV2-6"], state="Open",
                  sprints=["2026.09"]),
        raw_issue("DEV2-5", parents=["DEV2-4", "DEV2-99"], type_="Task", state="Verified", resolved=DONE_MS,
                  sprints=["2026.09"]),
        raw_issue("DEV2-6", parents=["DEV2-4"], type_="Task", state="In Progress", sprints=["2026.09"],
                  depends=[("DEV2-5", True), ("DEVSY-7", False)], summary="제목 접두어 없음"),
    ])
    tree = wbs.build_tree("DEV2-1", fetch)
    activities = {
        "DEV2-1": [state_act(SEP1_MS, "상태", "In Progress")],
        "DEV2-5": [state_act(SEP1_MS, "State", "In Progress")],
    }
    rows = wbs.build_rows(tree, activities)
    return tree, rows, wbs.compute_rollups(tree, rows, TODAY)


def typed_tree():
    """유형 색 픽스처 — Epic 루트 아래 Feature(Task·Bug), 표에 없는 유형, 빈 유형."""
    fetch = FakeFetcher([
        raw_issue("DEV2-1", children=["DEV2-2", "DEV2-5", "DEV2-6"], type_="Epic", sprints=["2026.09"]),
        raw_issue("DEV2-2", parents=["DEV2-1"], children=["DEV2-3", "DEV2-4"], sprints=["2026.09"]),
        raw_issue("DEV2-3", parents=["DEV2-2"], type_="Task", sprints=["2026.09"]),
        raw_issue("DEV2-4", parents=["DEV2-2"], type_="Bug", sprints=["2026.10"]),
        raw_issue("DEV2-5", parents=["DEV2-1"], type_="Spike", sprints=["2026.10"]),
        raw_issue("DEV2-6", parents=["DEV2-1"], type_=None),
    ])
    tree = wbs.build_tree("DEV2-1", fetch)
    rows = wbs.build_rows(tree, {})
    return tree, rows, wbs.compute_rollups(tree, rows, TODAY)


class TreeTraversalTests(unittest.TestCase):
    def test_depth_two_tree_orders_siblings_by_planned_start_then_id(self) -> None:
        fetch = FakeFetcher([
            raw_issue("DEV2-1", children=["DEV2-3", "DEVSY-9", "DEV2-2"], type_="Epic"),
            raw_issue("DEV2-2", sprints=["2026.09"], parents=["DEV2-1"]),
            raw_issue("DEV2-3", sprints=["2026.08"], children=["DEV2-4"], parents=["DEV2-1"]),
            raw_issue("DEVSY-9", parents=["DEV2-1"]),
            raw_issue("DEV2-4", type_="Task", parents=["DEV2-3"]),
        ])

        tree = wbs.build_tree("DEV2-1", fetch)

        self.assertEqual(tree.order, ["DEV2-1", "DEV2-3", "DEV2-4", "DEV2-2", "DEVSY-9"])
        self.assertEqual(tree.wbs["DEV2-4"], "1.1.1")
        self.assertEqual(tree.wbs["DEVSY-9"], "1.3")
        self.assertEqual(wbs.depth_distribution(tree), {1: 1, 2: 3, 3: 1})
        self.assertEqual(len(fetch.calls), 3)  # 레벨마다 한 번에 조회

    def test_depth_four_tree_epic_epic_feature_task(self) -> None:
        fetch = FakeFetcher([
            raw_issue("DEV2-1", children=["DEV2-10"], type_="Epic"),
            raw_issue("DEV2-10", children=["DEV2-20"], type_="Epic", parents=["DEV2-1"]),
            raw_issue("DEV2-20", children=["DEV2-31", "DEV2-30"], parents=["DEV2-10"]),
            raw_issue("DEV2-30", type_="Task", parents=["DEV2-20"]),
            raw_issue("DEV2-31", type_="Task", parents=["DEV2-20"]),
        ])

        tree = wbs.build_tree("DEV2-1", fetch)

        self.assertEqual(tree.wbs["DEV2-31"], "1.1.1.2")
        self.assertEqual(tree.depth["DEV2-31"], 4)
        self.assertEqual(tree.multi_parents, {})
        plan = wbs.outline_plan(tree, collapse=True)
        self.assertEqual([plan[i]["level"] for i in tree.order], [0, 1, 2, 3, 3])
        self.assertTrue(plan["DEV2-20"]["collapsed"])  # 하위가 전부 Task → 접힘
        self.assertFalse(plan["DEV2-10"]["collapsed"])
        self.assertTrue(plan["DEV2-30"]["hidden"])

    def test_cycle_is_expanded_once_and_recorded(self) -> None:
        fetch = FakeFetcher([
            raw_issue("DEV2-1", children=["DEV2-2"]),
            raw_issue("DEV2-2", children=["DEV2-3"], parents=["DEV2-1"]),
            raw_issue("DEV2-3", children=["DEV2-1"], parents=["DEV2-2"]),
        ])

        tree = wbs.build_tree("DEV2-1", fetch)

        self.assertEqual(tree.order, ["DEV2-1", "DEV2-2", "DEV2-3"])
        self.assertEqual(tree.cycles, [("DEV2-3", "DEV2-1")])
        fetched = [i for call in fetch.calls for i in call]
        self.assertEqual(len(fetched), len(set(fetched)))

    def test_multi_parent_listed_once_and_recorded(self) -> None:
        fetch = FakeFetcher([
            raw_issue("DEV2-1", children=["DEV2-2", "DEV2-3"]),
            raw_issue("DEV2-2", children=["DEV2-9"], parents=["DEV2-1"]),
            raw_issue("DEV2-3", children=["DEV2-9"], parents=["DEV2-1"]),
            raw_issue("DEV2-9", parents=["DEV2-2", "DEV2-3"]),
        ])

        tree = wbs.build_tree("DEV2-1", fetch)

        self.assertEqual(tree.order.count("DEV2-9"), 1)
        self.assertEqual(tree.parent["DEV2-9"], "DEV2-2")
        self.assertEqual(tree.multi_parents, {"DEV2-9": ["DEV2-2", "DEV2-3"]})

    def test_missing_child_is_skipped_and_recorded(self) -> None:
        fetch = FakeFetcher([raw_issue("DEV2-1", children=["DEV2-2", "DEVDG-404"]),
                             raw_issue("DEV2-2", parents=["DEV2-1"])])

        tree = wbs.build_tree("DEV2-1", fetch)

        self.assertEqual(tree.missing, ["DEVDG-404"])
        self.assertEqual(tree.order, ["DEV2-1", "DEV2-2"])


class ScheduleTests(unittest.TestCase):
    def test_planned_dates_span_min_to_max_sprint_month(self) -> None:
        rec = wbs.normalize_issue(raw_issue("DEV2-1", sprints=["2026.09", "2026.07", "2027.02"]))
        self.assertEqual(wbs.planned_dates(rec), (dt.date(2026, 7, 1), dt.date(2027, 2, 28)))

    def test_planned_dates_ignore_unparseable_and_empty(self) -> None:
        self.assertEqual(wbs.planned_dates({"sprints": ["Backlog", "2026.12"]}),
                         (dt.date(2026, 12, 1), dt.date(2026, 12, 31)))
        self.assertEqual(wbs.planned_dates({"sprints": []}), (None, None))
        self.assertEqual(wbs.planned_dates({"sprints": ["2026.13"]}), (None, None))

    def test_actual_start_accepts_localized_and_english_state_field(self) -> None:
        sep10 = DONE_MS
        sep20 = sep10 + 10 * 86400 * 1000
        localized = [state_act(sep20, "상태", "In Progress"), state_act(sep10, "상태", "Open")]
        english = [state_act(sep20, "State", "In Progress"), state_act(sep10, "State", "In Progress")]
        self.assertEqual(wbs.actual_start(localized), dt.date(2026, 9, 20))
        self.assertEqual(wbs.actual_start(english), dt.date(2026, 9, 10))

    def test_actual_start_none_without_in_progress_history(self) -> None:
        acts = [state_act(DONE_MS, "상태", "Fixed"), state_act(DONE_MS, "Sprints", "In Progress"),
                {"timestamp": DONE_MS, "field": {"name": "상태"}, "added": {"name": "Open"}}]
        self.assertIsNone(wbs.actual_start(acts))
        self.assertIsNone(wbs.actual_start(None))

    def test_actual_interval_extends_ongoing_work_to_today(self) -> None:
        row = {"act_start": dt.date(2026, 9, 1), "act_end": None}
        self.assertEqual(wbs.actual_interval(row, TODAY), (dt.date(2026, 9, 1), TODAY))
        row = {"act_start": None, "act_end": dt.date(2026, 9, 10)}
        self.assertEqual(wbs.actual_interval(row, TODAY), (dt.date(2026, 9, 10), dt.date(2026, 9, 10)))


class TimelineTests(unittest.TestCase):
    def test_day_columns_cover_each_date_and_week_columns_start_monday(self) -> None:
        start, end = dt.date(2026, 8, 1), dt.date(2026, 10, 31)
        days = wbs.timeline_columns(start, end, "day")
        weeks = wbs.timeline_columns(start, end, "week")

        self.assertEqual(len(days), 92)
        self.assertEqual(days[0], (start, start))
        self.assertEqual(days[-1], (end, end))
        self.assertEqual(len(weeks), 14)
        self.assertEqual(weeks[0], (dt.date(2026, 7, 27), dt.date(2026, 8, 2)))
        self.assertTrue(all(c0.weekday() == 0 for c0, _ in weeks))

    def test_weekend_detection(self) -> None:
        self.assertTrue(wbs.is_weekend(dt.date(2026, 10, 10)))  # 토
        self.assertTrue(wbs.is_weekend(dt.date(2026, 10, 11)))  # 일
        self.assertFalse(wbs.is_weekend(TODAY))  # 수

    def test_cell_kind_actual_wins_over_plan_and_cancel_is_grey(self) -> None:
        tree, rows, rollups = sample_tree()
        by_id = {r["id"]: r for r in rows}

        def kind(iid, day):
            r = by_id[iid]
            return wbs.cell_kind(r, wbs.actual_interval(r, TODAY), rollups.get(iid), day, day)

        self.assertEqual(kind("DEV2-1", dt.date(2026, 8, 10)), "plan")
        self.assertEqual(kind("DEV2-1", dt.date(2026, 9, 5)), "act_in")
        self.assertEqual(kind("DEV2-2", dt.date(2026, 9, 10)), "act_out")
        self.assertEqual(kind("DEV2-4", dt.date(2026, 9, 5)), "ru_act")  # 하위 실제가 자기 계획보다 우선
        self.assertEqual(kind("DEV2-4", dt.date(2026, 9, 20)), "plan")
        self.assertEqual(kind("DEV2-3", dt.date(2026, 9, 10)), "cancel_act")
        self.assertIsNone(kind("DEV2-3", dt.date(2026, 9, 11)))


class ProgressRollupCheckTests(unittest.TestCase):
    def test_progress_excludes_cancelled_leaves_from_denominator(self) -> None:
        _, rows, _ = sample_tree()
        by_id = {r["id"]: r for r in rows}

        self.assertEqual((by_id["DEV2-4"]["leaf_done"], by_id["DEV2-4"]["leaf_total"]), (1, 2))
        self.assertEqual((by_id["DEV2-1"]["leaf_done"], by_id["DEV2-1"]["leaf_total"]), (2, 3))
        self.assertAlmostEqual(by_id["DEV2-1"]["progress"], 2 / 3)
        self.assertIsNone(by_id["DEV2-3"]["progress"])  # 취소 leaf 자신은 대상 없음
        self.assertTrue(by_id["DEV2-3"]["cancelled"])
        self.assertEqual(by_id["DEV2-6"]["bucket"], wbs.BUCKET_DOING)
        self.assertEqual(by_id["DEV2-4"]["bucket"], wbs.BUCKET_TODO)

    def test_rollup_applies_when_children_are_wider_or_own_schedule_missing(self) -> None:
        _, _, rollups = sample_tree()

        self.assertFalse(rollups["DEV2-4"]["plan_applied"])  # 자기 Sprints와 하위 범위가 같음
        self.assertTrue(rollups["DEV2-4"]["act_applied"])  # 자기 실제 없음, 하위는 9/1~9/10
        self.assertEqual(rollups["DEV2-4"]["act"], (dt.date(2026, 9, 1), dt.date(2026, 9, 10)))
        self.assertFalse(rollups["DEV2-1"]["plan_applied"])
        self.assertFalse(rollups["DEV2-1"]["act_applied"])
        self.assertEqual(wbs.rollup_applied_count(rollups), 1)

    def test_rollup_ignores_cancelled_children_and_fills_missing_sprints(self) -> None:
        fetch = FakeFetcher([
            raw_issue("DEV2-1", children=["DEV2-2", "DEV2-3", "DEV2-4"], type_="Epic"),
            raw_issue("DEV2-2", parents=["DEV2-1"], sprints=["2026.08"]),
            raw_issue("DEV2-3", parents=["DEV2-1"], sprints=["2026.11"]),
            raw_issue("DEV2-4", parents=["DEV2-1"], sprints=["2027.01"], state="Duplicate", resolved=DONE_MS),
        ])
        tree = wbs.build_tree("DEV2-1", fetch)
        rows = wbs.build_rows(tree, {})

        rollups = wbs.compute_rollups(tree, rows, TODAY)

        self.assertTrue(rollups["DEV2-1"]["plan_applied"])
        self.assertEqual(rollups["DEV2-1"]["plan"], (dt.date(2026, 8, 1), dt.date(2026, 11, 30)))

    def test_checks_flag_structure_and_schedule_gaps(self) -> None:
        tree, rows, _ = sample_tree()

        checks = {c["key"]: c for c in wbs.compute_checks(tree, rows, TODAY)}
        ids = {k: [i for i, _ in c["items"]] for k, c in checks.items()}

        self.assertEqual(ids["root_flat"], ["DEV2-2", "DEV2-3"])
        self.assertIn("루트 직속 자식 3건 중 leaf 2건", checks["root_flat"]["note"])
        self.assertEqual(sorted(ids["prefix_variant"]), ["DEV2-1", "DEV2-3", "DEV2-6"])
        self.assertEqual(ids["no_sprints"], ["DEV2-3"])
        self.assertEqual(ids["done_without_progress"], ["DEV2-2"])  # 취소성 종료 DEV2-3 제외
        self.assertEqual(ids["doing_without_progress"], ["DEV2-6"])
        self.assertEqual(ids["blocked_in_progress"], ["DEV2-6"])
        self.assertIn("DEVSY-7", checks["blocked_in_progress"]["items"][0][1])
        self.assertEqual(ids["multi_parent"], ["DEV2-5"])
        self.assertEqual(ids["overdue"], ["DEV2-4", "DEV2-6"])  # 루트는 2026.10까지라 제외
        self.assertEqual(ids["cancelled"], ["DEV2-3"])
        self.assertEqual(ids["cycle"], [])
        self.assertEqual(ids["missing"], [])

    def test_team_summary_total_row_excludes_cancelled_from_progress(self) -> None:
        _, rows, _ = sample_tree()
        summary = {s["project"]: s for s in wbs.team_summary(rows)}
        self.assertEqual(summary["전체"]["count"], 6)
        self.assertEqual(summary["전체"]["cancelled"], 1)
        self.assertEqual((summary["전체"]["leaf_done"], summary["전체"]["leaves"]), (2, 3))
        self.assertAlmostEqual(summary["전체"]["progress"], 2 / 3)
        self.assertEqual(summary["DEV2"]["no_plan"], 1)
        self.assertEqual(summary["전체"]["no_actual"], 2)  # 점검과 같은 기준 — 취소성 종료 제외


class WorkbookTests(unittest.TestCase):
    def write(self, unit):
        from openpyxl import load_workbook

        tree, rows, rollups = sample_tree()
        checks = wbs.compute_checks(tree, rows, TODAY)
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "wbs.xlsx"
            wbs.write_workbook(out, tree, rows, checks, rollups, [("루트 이슈", "DEV2-1")], BASE, TODAY, unit)
            return tree, load_workbook(out)

    def test_day_gantt_has_one_column_per_date_and_outline_groups(self) -> None:
        tree, book = self.write("day")

        self.assertEqual(book.sheetnames, ["WBS", "간트", "팀별 요약", "점검", "메타"])
        gantt = book["간트"]
        self.assertEqual(gantt.max_column, 6 + 92)  # 2026-08-01 ~ 2026-10-31
        self.assertEqual(gantt.max_row, 5 + len(tree.order))
        self.assertFalse(gantt.sheet_properties.outlinePr.summaryBelow)
        task_rows = [6 + tree.order.index(i) for i in ("DEV2-5", "DEV2-6")]
        for row_no in task_rows:
            self.assertTrue(gantt.row_dimensions[row_no].hidden)
            self.assertGreaterEqual(gantt.row_dimensions[row_no].outline_level, 2)
        self.assertEqual(gantt.cell(row=6, column=1).fill.fgColor.rgb, "001F3864")  # 레벨 1
        self.assertEqual(gantt.cell(row=7, column=1).fill.fgColor.rgb, "00D9E2F3")  # 레벨 2
        sat = 6 + 1  # 2026-08-01(토)이 첫 날짜 열
        self.assertEqual(gantt.cell(row=8, column=sat).fill.fgColor.rgb, "00F2F2F2")

    def test_week_gantt_and_wbs_sheet_rows(self) -> None:
        tree, book = self.write("week")

        self.assertEqual(book["간트"].max_column, 6 + 14)
        sheet = book["WBS"]
        self.assertEqual(sheet.max_row, len(tree.order) + 1)
        self.assertEqual(sheet["C2"].value, "DEV2-1")
        self.assertEqual(sheet["C2"].hyperlink.target, f"{BASE}/issue/DEV2-1")
        self.assertEqual(sheet["D3"].alignment.indent, 1)
        self.assertFalse(sheet.row_dimensions[2 + tree.order.index("DEV2-5")].hidden)  # WBS는 전체 펼침
        self.assertEqual(sheet.row_dimensions[2 + tree.order.index("DEV2-5")].outline_level, 2)

    def test_output_inside_repo_is_rejected(self) -> None:
        with self.assertRaises(ValueError):
            wbs.ensure_outside_repo(wbs.REPO / "wbs.xlsx")
        self.assertEqual(wbs.ensure_outside_repo(Path("/private/tmp/wbs.xlsx")).name, "wbs.xlsx")


class HtmlTests(unittest.TestCase):
    def payload(self):
        tree, rows, rollups = sample_tree()
        checks = wbs.compute_checks(tree, rows, TODAY)
        meta = {"root": "DEV2-1", "root_title": "루트", "extracted": "2026-10-07 09:00:00 KST", "api_calls": 3}
        return rows, wbs.build_html_payload(tree, rows, checks, rollups, BASE, TODAY, meta)

    def test_payload_shares_progress_and_rollup_with_workbook(self) -> None:
        rows, payload = self.payload()
        by_id = {r["id"]: r for r in payload["rows"]}

        self.assertEqual(len(payload["rows"]), len(rows))
        self.assertAlmostEqual(payload["kpi"]["progress"], wbs.team_summary(rows)[-1]["progress"])
        self.assertAlmostEqual(payload["kpi"]["progress"], 2 / 3)
        self.assertEqual((payload["kpi"]["done"], payload["kpi"]["cancelled"]), (2, 1))
        self.assertEqual(payload["meta"]["rollups"], 1)
        self.assertEqual(by_id["DEV2-4"]["ra"], ["2026-09-01", "2026-09-10"])
        self.assertIsNone(by_id["DEV2-4"]["rp"])
        self.assertEqual(by_id["DEV2-3"]["b"], wbs.BUCKET_CANCEL)
        self.assertEqual(payload["range"], ["2026-08-01", "2026-10-31"])

    def test_render_inlines_escaped_json_without_external_scripts(self) -> None:
        _, payload = self.payload()
        payload["rows"][0]["title"] = "</script><b>위험</b>"

        page = wbs.render_html(payload, "SF 프로젝트 간트")

        self.assertIn("<title>SF 프로젝트 간트</title>", page)
        self.assertNotIn("</script><b>", page)
        data = re.search(r'<script type="application/json" id="wbs-data">(.*?)</script>', page, re.S).group(1)
        parsed = json.loads(data)
        self.assertEqual(len(parsed["rows"]), 6)
        self.assertEqual(parsed["rows"][0]["title"], "</script><b>위험</b>")
        self.assertNotIn("<script src", page)
        self.assertNotIn("<link", page)
        svg_ns = "http://www.w3.org/2000/svg"  # createElementNS 네임스페이스 — 네트워크 요청 아님
        self.assertEqual(set(re.findall(r"https?://[^\s\"'<>\\)]+", page)) - {svg_ns}, {BASE})

    def test_html_path_inside_repo_is_rejected_before_any_request(self) -> None:
        code = wbs.main(["DEV2-1", "--out", "/private/tmp/wbs.xlsx", "--html", str(wbs.REPO / "gantt.html")])
        self.assertEqual(code, 2)


class DateLabelAndLinkTests(unittest.TestCase):
    def test_range_label_formats(self) -> None:
        d = dt.date
        self.assertEqual(wbs.fmt_md(d(2026, 4, 6)), "4/6")
        self.assertEqual(wbs.range_label(d(2026, 4, 1), d(2026, 5, 31)), "4/1–5/31")
        self.assertEqual(wbs.range_label(d(2026, 4, 13), None, ongoing=True), "4/13–진행")
        self.assertEqual(wbs.range_label(None, d(2026, 5, 7)), "?–5/7")
        self.assertEqual(wbs.range_label(None, None), "")

    def test_late_days_counts_beyond_plan_end(self) -> None:
        row = {"plan_end": dt.date(2026, 9, 30), "act_start": dt.date(2026, 9, 1), "act_end": dt.date(2026, 10, 12)}
        self.assertEqual(wbs.late_days(row, TODAY), 12)
        row = {"plan_end": dt.date(2026, 9, 30), "act_start": dt.date(2026, 9, 1), "act_end": None}
        self.assertEqual(wbs.late_days(row, TODAY), 7)  # 진행 중이면 기준일까지
        row = {"plan_end": dt.date(2026, 9, 30), "act_start": None, "act_end": dt.date(2026, 9, 20)}
        self.assertIsNone(wbs.late_days(row, TODAY))

    def test_collect_links_splits_depend_and_relates_inside_and_outside_tree(self) -> None:
        fetch = FakeFetcher([
            raw_issue("DEV2-1", children=["DEV2-2", "DEV2-3"], type_="Epic"),
            raw_issue("DEV2-2", parents=["DEV2-1"], required_for=["DEV2-3", "DEV2-900"], relates=["DEV2-3"]),
            raw_issue("DEV2-3", parents=["DEV2-1"], depends=[("DEV2-2", False), ("DEVSY-50", True)],
                      relates=["DEV2-2", "DEVDG-7"]),
        ])
        tree = wbs.build_tree("DEV2-1", fetch)

        links = wbs.collect_links(tree)
        pairs = {(l["kind"], l["src"], l["dst"], l["in_tree"]) for l in links}

        self.assertEqual(pairs, {
            ("Depend", "DEV2-2", "DEV2-3", True),  # 양쪽에서 보이지만 한 번만
            ("Depend", "DEV2-2", "DEV2-900", False),
            ("Depend", "DEVSY-50", "DEV2-3", False),
            ("Relates", "DEV2-2", "DEV2-3", True),
            ("Relates", "DEV2-3", "DEVDG-7", False),
        })
        self.assertEqual(wbs.link_counts(links), {"Depend": {"in_tree": 1, "external": 2},
                                                  "Relates": {"in_tree": 1, "external": 1}})

    def test_violation_when_successor_starts_before_predecessor_ends(self) -> None:
        d = dt.date
        pred = {"act_start": d(2026, 9, 1), "act_end": d(2026, 9, 20), "plan_start": None, "plan_end": None}
        early = {"act_start": d(2026, 9, 10), "act_end": None, "plan_start": d(2026, 10, 1), "plan_end": None}
        after = {"act_start": None, "act_end": None, "plan_start": d(2026, 9, 20), "plan_end": None}
        unknown = {"act_start": None, "act_end": None, "plan_start": None, "plan_end": None}
        self.assertTrue(wbs.is_violation(pred, early))  # 실제 시작 우선
        self.assertFalse(wbs.is_violation(pred, after))  # 같은 날 시작은 위반 아님
        self.assertFalse(wbs.is_violation(pred, unknown))

    def test_violation_check_and_html_links_agree(self) -> None:
        fetch = FakeFetcher([
            raw_issue("DEV2-1", children=["DEV2-2", "DEV2-3"], type_="Epic"),
            raw_issue("DEV2-2", parents=["DEV2-1"], sprints=["2026.09"], required_for=["DEV2-3"]),
            raw_issue("DEV2-3", parents=["DEV2-1"], sprints=["2026.09"], depends=[("DEV2-2", False)]),
        ])
        tree = wbs.build_tree("DEV2-1", fetch)
        rows = wbs.build_rows(tree, {})
        rollups = wbs.compute_rollups(tree, rows, TODAY)
        checks = {c["key"]: c for c in wbs.compute_checks(tree, rows, TODAY)}

        self.assertEqual([i for i, _ in checks["dependency_violation"]["items"]], ["DEV2-3"])
        payload = wbs.build_html_payload(tree, rows, list(checks.values()), rollups, BASE, TODAY,
                                         {"root": "DEV2-1", "extracted": "-", "api_calls": 0})
        self.assertEqual(payload["links"], [{"k": "d", "f": "DEV2-2", "t": "DEV2-3", "v": True}])


class HtmlParityTests(unittest.TestCase):
    def test_html_rows_carry_every_excel_wbs_column(self) -> None:
        from openpyxl import load_workbook

        tree, rows, rollups = sample_tree()
        checks = wbs.compute_checks(tree, rows, TODAY)
        payload = wbs.build_html_payload(tree, rows, checks, rollups, BASE, TODAY,
                                         {"root": "DEV2-1", "extracted": "-", "api_calls": 0})
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "wbs.xlsx"
            wbs.write_workbook(out, tree, rows, checks, rollups, [], BASE, TODAY)
            header = [c.value for c in load_workbook(out)["WBS"][1]]

        self.assertEqual(header, wbs.WBS_LABELS)
        self.assertEqual(set(wbs.HTML_KEYS), set(wbs.WBS_LABELS))
        for row in payload["rows"]:
            self.assertTrue(all(wbs.HTML_KEYS[label] in row for label in wbs.WBS_LABELS), row["id"])

    def test_html_summary_and_check_counts_equal_workbook_functions(self) -> None:
        tree, rows, rollups = sample_tree()
        checks = wbs.compute_checks(tree, rows, TODAY)
        payload = wbs.build_html_payload(tree, rows, checks, rollups, BASE, TODAY,
                                         {"root": "DEV2-1", "extracted": "-", "api_calls": 0})

        excel_summary = wbs.team_summary(rows)
        self.assertEqual([{k: v for k, v in s.items() if k != "label"} for s in payload["summary"]], excel_summary)
        self.assertEqual({c["key"]: len(c["items"]) for c in payload["checks"]},
                         {c["key"]: len(c["items"]) for c in checks})
        self.assertIn("dependency_violation", payload["pinned"])

    def test_template_fits_width_and_uses_column_width_variables(self) -> None:
        page = wbs.HTML_TEMPLATE.read_text(encoding="utf-8")

        self.assertIn("function fitDw()", page)
        self.assertIn("ResizeObserver", page)
        self.assertIn('setProperty("--dw"', page)
        self.assertNotIn("--dw: 4px", page)  # 고정 일폭에 의존하지 않는다
        self.assertIn("--col-title", page)
        self.assertIn('setProperty("--col-" + c.k', page)
        self.assertIn('store("widths"', page)


class TypeColorTests(unittest.TestCase):
    EXPECTED = {"DEV2-1": "epic", "DEV2-2": "feature", "DEV2-3": "task", "DEV2-4": "bug",
                "DEV2-5": "other", "DEV2-6": "other"}

    def test_type_color_key_maps_known_types_and_falls_back(self) -> None:
        self.assertEqual(wbs.type_color_key("Epic"), "epic")
        self.assertEqual(wbs.type_color_key(" bug "), "bug")
        for name in ("Spike", "", None):
            self.assertEqual(wbs.type_color_key(name), wbs.TYPE_COLOR_FALLBACK)
        self.assertIn(wbs.TYPE_COLOR_FALLBACK, wbs.TYPE_COLORS)

    def test_wbs_sheet_type_cell_keeps_type_fill_over_level_fill(self) -> None:
        from openpyxl import load_workbook

        tree, rows, rollups = typed_tree()
        checks = wbs.compute_checks(tree, rows, TODAY)
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "wbs.xlsx"
            wbs.write_workbook(out, tree, rows, checks, rollups, [], BASE, TODAY)
            book = load_workbook(out)
        sheet = book["WBS"]
        row_of = {sheet.cell(row=i, column=3).value: i for i in range(2, sheet.max_row + 1)}

        self.assertEqual(sheet.cell(row=1, column=5).value, "유형")
        for iid, key in self.EXPECTED.items():
            cell = sheet.cell(row=row_of[iid], column=5)
            self.assertEqual(cell.fill.fgColor.rgb, "00" + wbs.TYPE_COLORS[key]["bg"], iid)
            self.assertEqual(cell.font.color.rgb, "00" + wbs.TYPE_COLORS[key]["ink"], iid)
        self.assertEqual(sheet.cell(row=row_of["DEV2-1"], column=4).fill.fgColor.rgb, "001F3864")  # 레벨 1 채움 유지
        self.assertEqual(sheet.cell(row=row_of["DEV2-2"], column=4).fill.fgColor.rgb, "00D9E2F3")  # 레벨 2 채움 유지
        self.assertTrue(sheet.cell(row=row_of["DEV2-1"], column=5).font.bold)
        gantt = book["간트"]  # 간트 시트에는 유형 칸이 없다 — 레벨 채움 그대로
        self.assertEqual(gantt.cell(row=6, column=5).fill.fgColor.rgb, "001F3864")

    def test_html_payload_and_template_carry_type_colors(self) -> None:
        tree, rows, rollups = typed_tree()
        checks = wbs.compute_checks(tree, rows, TODAY)
        payload = wbs.build_html_payload(tree, rows, checks, rollups, BASE, TODAY,
                                         {"root": "DEV2-1", "extracted": "-", "api_calls": 0})

        self.assertEqual({r["id"]: r["tc"] for r in payload["rows"]}, self.EXPECTED)
        self.assertEqual(set(payload["type_colors"]), set(wbs.TYPE_COLORS))
        for key, colors in wbs.TYPE_COLORS.items():
            self.assertEqual(payload["type_colors"][key], {k: "#" + v for k, v in colors.items()})
        page = wbs.render_html(payload, "유형 색")
        data = re.search(r'<script type="application/json" id="wbs-data">(.*?)</script>', page, re.S).group(1)
        self.assertEqual(json.loads(data)["type_colors"], payload["type_colors"])
        self.assertIn('"tybdg ty-" + r.tc', page)  # 표 '유형' 칸 배지
        self.assertIn("TC[k].bg_dark", page)  # 다크 테마 토큰도 같은 표에서
        self.assertEqual(page.count("--ty-ring:"), 3)  # 라이트 + 다크 두 블록


if __name__ == "__main__":
    unittest.main()
